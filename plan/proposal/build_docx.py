"""proposal-v0.md → kaynak/CSE491_Project_Proposal_Template.docx şablonuyla doldurulmuş .docx
Çalıştır: python3 plan/proposal/build_docx.py
Kaynak metin tek yer: plan/proposal/proposal-v0.md (burada elle metin yazma).
"""
import copy, re, pathlib
import docx

KOK = pathlib.Path(__file__).resolve().parents[2]
MD = KOK / "plan/proposal/proposal-v0.md"
TPL = KOK / "kaynak/CSE491_Project_Proposal_Template.docx"
OUT = KOK / "plan/proposal/NutriScan_CSE491_Proposal_v0.docx"

# ---------- markdown'ı bölümlere ayır
text = MD.read_text(encoding="utf-8")
text = text.split("\n---\n", 1)[1] if text.startswith("---") else text
sections = {}
cur = "_head"
for line in text.splitlines():
    m = re.match(r"^## (\d)\. ", line)
    if m:
        cur = m.group(1); sections[cur] = []; continue
    sections.setdefault(cur, []).append(line)

def blocks(lines):
    """satırları bloklara ayırır: ('p', metin) | ('label', metin) | ('table', rows) | ('li', metin)"""
    out, tbl = [], []
    for l in lines + [""]:
        if l.startswith("|"):
            cells = [c.strip() for c in l.strip().strip("|").split("|")]
            if not all(re.fullmatch(r"-+", c) for c in cells if c):
                tbl.append(cells)
            continue
        if tbl:
            out.append(("table", tbl)); tbl = []
        s = l.strip()
        if not s:
            continue
        m = re.fullmatch(r"\*\*(.+?)\*\*", s)
        if m:
            out.append(("label", m.group(1))); continue
        if re.match(r"^(\d+\.|-)\s", s):
            out.append(("li", re.sub(r"^-\s", "• ", s))); continue
        out.append(("p", s))
    return out

def plain(s):
    return s.replace("**", "").replace("`", "")

def add_runs(par, s):
    """**kalın** destekli run ekleme"""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", s)):
        if part:
            r = par.add_run(part.replace("`", ""))
            if i % 2 == 1:
                r.bold = True

# ---------- şablon
d = docx.Document(TPL)
P = d.paragraphs

def insert_after(anchor, s, style=None):
    new = copy.deepcopy(anchor._p)
    for child in list(new):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            new.remove(child)
    anchor._p.addnext(new)
    par = docx.text.paragraph.Paragraph(new, anchor._parent)
    add_runs(par, s)
    return par

def delete(par):
    par._p.getparent().remove(par._p)

def fill_table(t, rows, start=1):
    """rows: başlık hariç satırlar; gerekirse son satırı kopyalayarak ekler"""
    while len(t.rows) - start < len(rows):
        t._tbl.append(copy.deepcopy(t.rows[-1]._tr))
    while len(t.rows) - start > len(rows):
        t._tbl.remove(t.rows[-1]._tr)
    for ri, row in enumerate(rows):
        cells = t.rows[start + ri].cells
        for ci, val in enumerate(row[: len(cells)]):
            cell = cells[ci]
            for p in cell.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            p0 = cell.paragraphs[0]
            for r in list(p0.runs):
                r._r.getparent().remove(r._r)
            add_runs(p0, val)

# başlık tablosu ve üyeler
head = blocks(sections["_head"])
tables = [b[1] for b in head if b[0] == "table"]
kv = {r[0]: r[1] for r in tables[0] if len(r) >= 2}
t0 = d.tables[0]
for row in t0.rows:
    k = row.cells[0].text.strip()
    if k in kv:
        c = row.cells[1]; p0 = c.paragraphs[0]
        for r in list(p0.runs): r._r.getparent().remove(r._r)
        add_runs(p0, kv[k])
fill_table(d.tables[1], tables[1][1:])

# üstteki talimat paragrafı (2) sil
delete(P[2])

def fill_section(heading_text, sec_key, table_map=None):
    """başlıktan sonraki italik talimatı ve boş paragrafları siler, bloklarla doldurur.
    table_map: {sıra: tablo_index} — md'deki n. tabloyu şablondaki tabloya yazar."""
    paras = d.paragraphs
    hi = next(i for i, p in enumerate(paras) if p.text.strip().startswith(heading_text))
    # başlıktan sonraki bölüm paragrafları (bir sonraki Heading 1'e kadar)
    j = hi + 1
    body = []
    while j < len(paras) and paras[j].style.name != "Heading 1":
        body.append(paras[j]); j += 1
    instr = body[0]
    labels = {p.text.strip(): p for p in body[1:] if p.text.strip()}
    for p in body[1:]:
        if not p.text.strip():
            delete(p)
    bl = blocks(sections[sec_key])
    anchor = instr
    tcount = 0
    pending_empty = []
    for kind, val in bl:
        if kind == "label" and val in labels:
            anchor = labels[val]; continue
        if kind == "label":
            anchor = insert_after(anchor, f"**{val}**"); continue
        if kind == "table":
            if table_map and tcount in table_map:
                t = d.tables[table_map[tcount]]
                fill_table(t, val[1:])
                # tablodan sonra gelen bloklar tablonun arkasına yazılsın
                np_ = copy.deepcopy(instr._p)
                for child in list(np_):
                    if child.tag.endswith("}r"):
                        np_.remove(child)
                t._tbl.addnext(np_)
                anchor = docx.text.paragraph.Paragraph(np_, instr._parent)
                pending_empty.append(anchor)
            tcount += 1; continue
        anchor = insert_after(anchor, val)
    delete(instr)
    for e in pending_empty:
        if not e.text.strip():
            delete(e)

fill_section("1. Summary", "1")
fill_section("2. Problem Definition", "2")
fill_section("3. Objectives", "3")
fill_section("4. Similar Systems", "4", {0: 2})
fill_section("5. Preliminary Requirements", "5", {0: 3, 1: 4})
fill_section("6. Method and Technology Plan", "6", {0: 5})
fill_section("7. Work Plan", "7", {0: 6})
fill_section("8. Expected Outputs", "8", {0: 7})
fill_section("9. References", "9")

# tabloları sayfa genişliğine (18 cm) yay — dar sütunlar satırları uzatıp 6 sayfa sınırını aşıyordu
from docx.oxml.ns import qn
def set_widths(t, cms):
    tw = [str(int(c * 567)) for c in cms]
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), tw):
        gc.set(qn("w:w"), w)
    tblW = t._tbl.tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), str(sum(int(w) for w in tw))); tblW.set(qn("w:type"), "dxa")
    for row in t.rows:
        for tc, w in zip(row._tr.findall(qn("w:tc")), tw):
            tcW = tc.get_or_add_tcPr().get_or_add_tcW()
            tcW.set(qn("w:w"), w); tcW.set(qn("w:type"), "dxa")

WIDTHS = {2: [3.6, 5.6, 8.8], 3: [1.5, 12.6, 3.9], 4: [1.5, 10.6, 5.9],
          5: [3.2, 6.6, 8.2], 6: [1.3, 8.9, 5.6, 2.2], 7: [5.4, 2.1, 2.0, 8.5]}
for ti, cms in WIDTHS.items():
    set_widths(d.tables[ti], cms)

d.save(OUT)
print("yazıldı:", OUT)
